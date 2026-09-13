"""
API 端到端验收：导入 → 建库 → 分析 → 复核 → LLM → QC → 报告。

验证"完整可运行"这一交付要求，而不只是单元级规则正确。
"""
from __future__ import annotations

from fastapi.testclient import TestClient

from _harness import EI_TABLE, case, eq, ok
from pathlib import Path

_client = None
# 跨用例共享：API-04 建的分析与组分，供后续用例使用
SHARED = {"analysis_id": None, "component_id": "peak_007"}


def client():
    global _client
    if _client is None:
        from pmaa.api.main import app
        _client = TestClient(app)
    return _client


@case("API-01", "健康检查与枚举端点")
def _api01():
    c = client()
    h = c.get("/v1/health").json()
    eq(h["status"], "ok", "健康状态")
    e = c.get("/v1/enums").json()
    ok(len(e["qc_status"]) == 4, "qc_status 枚举")
    ok("LLM_DROPPED_AMBIGUITY" in e["error_codes"], "错误码清单")
    fw = c.get("/v1/chemistry/frameworks").json()
    eq(len(fw["items"]), 5, "五种骨架框架")


@case("API-02", "化学映射自检端点：Hexp 2,3,6 → 4-连接")
def _api02():
    c = client()
    r = c.post("/v1/chemistry/map", json={"framework_id": "aldohexopyranose",
                                          "o_methyl_positions": [2, 3, 6]}).json()
    eq(r["linkage_positions"], [4], "连接位")
    eq(r["o_acetyl_positions"], [1, 4, 5], "乙酰位")
    bad = c.post("/v1/chemistry/map", json={"framework_id": "aldohexopyranose",
                                           "o_methyl_positions": [2, 3, 5, 6]})
    eq(bad.status_code, 422, "不相容位点应报 422")
    eq(bad.json()["error_code"], "METHYL_POSITION_INCOMPATIBLE_WITH_FRAMEWORK", "稳定错误码")


@case("API-03", "导入 EI 汇总表并提交")
def _api03():
    if not Path(EI_TABLE).exists():
        print("        （跳过：EI 汇总表不可用）")
        return
    c = client()
    pv = c.post("/v1/imports/preview", json={"path": EI_TABLE}).json()
    ok(pv["summary"]["total_rows"] == 57, f"应识别 57 行，实际 {pv['summary']['total_rows']}")
    eq(pv["summary"]["detected_semantics"]["red_means"], "reported_gt20pct", "红色语义")
    cm = c.post("/v1/imports/commit", json={"path": EI_TABLE, "reviewer": "test"}).json()
    ok(cm["summary"]["imported"] >= 50, "入库条数")
    ok(cm["summary"]["isolated"] == 0, "不应有错误隔离行")
    refs = c.get("/v1/references").json()
    ok(refs["count"] >= 50, "参考条目数")


@case("API-04", "创建分析并返回五维状态与候选")
def _api04():
    c = client()
    c.post("/v1/samples", json={"sample_id": "api_sample", "code": "API-1"})
    c.post("/v1/methods", json={"method_id": "api_method", "ei_energy_ev": 70})
    r = c.post("/v1/analyses", json={
        "sample_id": "api_sample", "method_id": "api_method",
        "sample_reduction_label": "NaBD4",
        "components": [{
            "component_id": "peak_007", "rt": 23.4, "area": 12345.0,
            "evidence_type": "qualitative_ion_list",
            "ions": [{"mz": 118, "marker": "red"}, {"mz": 113, "marker": "+"},
                     {"mz": 99, "marker": "+"}, {"mz": 102, "marker": "+"},
                     {"mz": 233, "marker": "+"}]}]}).json()
    eq(r["status"], "succeeded", "分析状态")
    SHARED["analysis_id"] = r["analysis_id"]
    comp = r["components"][0]["result"]
    for f in ("analysis_id", "component_id", "input_hash", "method_version", "library_version",
              "rule_version", "qc_status", "derivative_identity_level", "residue_mapping_level",
              "quantification_level", "review_status", "candidate_derivatives",
              "candidate_residues", "positive_evidence", "conflicts", "missing_evidence",
              "match_scores", "quantification", "next_discriminating_evidence"):
        ok(f in comp, f"结果必须含字段 {f}")
    ok(len(comp["candidate_derivatives"]) > 0, "应有候选")
    eq(comp["derivative_identity_level"], "candidate", "定性表证据上限")
    eq(comp["quantification_level"], "area_only", "无校准只能给面积占比")


@case("API-04b", "实际证据包中 PMAA—残基关联与歧义 ID 一致")
def _api04b():
    from pmaa.analysis.evidence import build_bundle
    from pmaa.data import repositories as repo

    aid = SHARED["analysis_id"]
    analysis = repo.AnalysisRepo.get(aid)
    crow = repo.ComponentRepo.by_analysis(aid)[0]
    result = crow.result
    bundle = build_bundle(
        analysis_id=aid, analysis_revision=analysis.revision,
        component_id=crow.component_id, result=result, component_row=crow,
        library_version=analysis.library_version, rule_version=analysis.rule_version,
        method_version=analysis.method_id or "unknown_method",
        evidence_type=crow.evidence_type, spectrum_completeness=crow.spectrum_completeness,
        input_hash=crow.input_hash, qc_status=analysis.qc_status)
    ids = {x["candidate_id"] for x in bundle.candidates}
    for amb in bundle.ambiguities:
        ok(set(amb.candidate_ids) <= ids, "歧义成员必须引用真实候选 ID")
    residue_owner = {r["residue_label"]: r["derivative_id"]
                     for r in result["candidate_residues"]}
    for cand in bundle.candidates:
        if cand["residue_label"] != "unknown":
            eq(cand["derivative_id"], residue_owner[cand["residue_label"]],
               "不得生成跨 PMAA 的残基笛卡尔组合")


@case("API-05", "幂等键复用同一分析")
def _api05():
    c = client()
    body = {"idempotency_key": "idem_test_1", "components": [{
        "component_id": "p1", "evidence_type": "qualitative_ion_list",
        "ions": [{"mz": 118, "marker": "+"}]}]}
    a = c.post("/v1/analyses", json=body).json()
    b = c.post("/v1/analyses", json=body).json()
    eq(a["analysis_id"], b["analysis_id"], "幂等键应复用")
    eq(b["status"], "reused", "应标记复用")


@case("API-06", "人工复核必须给理由与证据")
def _api06():
    c = client()
    aid = SHARED["analysis_id"]
    bad = c.post(f"/v1/analyses/{aid}/reviews",
                 json={"component_id": "peak_007", "decision": {"derivative_identity": "confirmed"},
                       "reviewer": "t", "reason": ""})
    eq(bad.status_code, 422, "无理由应拒绝")
    bad2 = c.post(f"/v1/analyses/{aid}/reviews",
                  json={"component_id": "peak_007", "decision": {"derivative_identity": "confirmed"},
                        "reviewer": "t", "reason": "凭经验"})
    eq(bad2.status_code, 422, "无证据引用不得确认")
    good = c.post(f"/v1/analyses/{aid}/reviews",
                  json={"component_id": "peak_007",
                        "decision": {"derivative_identity": "probable", "evidence_ids": ["F001"]},
                        "reviewer": "t", "reason": "同方法标准保留一致"})
    ok(good.status_code == 200, "合法复核应通过")


@case("API-07", "LLM 评议端到端（Mock 适配器）")
def _api07():
    c = client()
    c.post("/v1/llm/providers", json={"provider_id": "api_mock", "name": "离线",
                                      "adapter": "mock", "model": "mock-deterministic-v1"})
    aid = SHARED["analysis_id"]
    r = c.post(f"/v1/analyses/{aid}/llm-assessments", json={
        "analysis_revision": 1, "component_ids": ["peak_007"],
        "provider_config_id": "api_mock"}).json()
    a = r["assessments"][0]
    eq(a["status"], "accepted", "Mock 输出应通过校验")
    got = c.get(f"/v1/llm-assessments/{a['assessment_id']}").json()
    ok(got["raw_response"], "必须保存原始响应")
    ok(got["prompt_hash"], "必须保存提示词哈希")
    ok(got["parsed_response"]["retained_ambiguity_ids"] is not None, "必须保留歧义")


@case("API-08", "QC 与报告导出")
def _api08():
    from pmaa.data import repositories as repo

    c = client()
    aid = SHARED["analysis_id"]
    q = c.post(f"/v1/analyses/{aid}/qc", json={"blank_peak_ratio": 0.01,
                                               "independent_preparations": 3,
                                               "repeat_cv": 0.08, "recovery": 0.95}).json()
    eq(q["qc_status"], "passed", "QC 应通过")
    rep = c.post(f"/v1/analyses/{aid}/reports", json={"fmt": "json"}).json()
    eq(rep["kind"], "formal", "正式报告")
    ok(rep["sha256"], "报告哈希")
    ok(rep["qc"] and rep["qc"]["checks"], "报告必须包含可追溯的 QC 明细")
    ok(rep["components"][0]["ions"], "报告必须包含峰谱/离子观测")
    ok(rep["components"][0]["review_history"], "报告必须包含人工复核轨迹")
    eq(rep["method"]["method_id"], "api_method", "报告必须冻结 GC 方法快照")
    eq(repo.MethodRepo.get("api_method").locked, True, "出报告后方法版本必须锁定")
    got_qc = c.get(f"/v1/analyses/{aid}/qc").json()
    eq(got_qc["payload"]["status"], "passed", "QC 明细必须持久化")
    lst = c.get(f"/v1/reports/{aid}").json()
    ok(len(lst["items"]) >= 1, "报告应归档")
    replay = c.get(f"/v1/reports/by-id/{rep['report_id']}").json()
    eq(replay["sha256"], rep["sha256"], "历史报告应原样回放")
    html = c.post(f"/v1/analyses/{aid}/reports", json={"fmt": "html"})
    eq(html.status_code, 200, "HTML 报告应通过 POST 导出")
    ok("text/html" in html.headers.get("content-type", ""), "HTML 响应类型")


@case("API-09", "来源链可追溯到源文件单元格")
def _api09():
    c = client()
    refs = c.get("/v1/references").json()["items"]
    p = c.get(f"/v1/references/{refs[0]['record_id']}/provenance").json()
    ok(p["source_file"]["sha256"], "必须能查到源文件哈希")
    ok(p["source_row_id"], "必须能查到源行")
    ok(len(p["ion_source_cells"]) > 0, "离子必须带源单元格")


@case("API-10", "错误行隔离与脱敏配置")
def _api10():
    c = client()
    r = c.get("/v1/llm/providers").json()["items"][0]
    ok("secret_ref" in r and "api_key" not in r, "配置端点不得回传密钥")
    t = c.post(f"/v1/llm/providers/{r['provider_id']}/test").json()
    ok("capabilities" in t, "能力测试应返回能力声明")
    ok(t["note"], "应声明不发送实验数据")


@case("API-11", "文件上传：支持格式矩阵与拒收规则")
def _api11():
    c = client()
    f = c.get("/v1/imports/formats").json()
    exts = set(f["supported_extensions"])
    for must in (".xlsx", ".csv", ".json", ".pdf", ".docx", ".png", ".jpg", ".tif"):
        ok(must in exts, f"应支持 {must}")
    ok(".zip" not in exts, "压缩包不应在支持列表")
    ok(f["max_upload_bytes"] > 0, "应声明体积上限")
    ok(".zip" in f["rejected"]["archives"], "压缩包应在拒收清单")


@case("API-12", "上传峰表（CSV）→ 解析 → 直接跑分析")
def _api12():
    c = client()
    csv_text = ("component_id,peak_label,rt,area,evidence_type,ions\n"
                "peak_101,P1,21.3,12000,qualitative_ion_list,\"118:red,129:+,87:+\"\n"
                "peak_102,P2,24.8,8000,qualitative_ion_list,\"102:+,118:red,162:+\"\n")
    r = c.post("/v1/imports/upload",
               files={"file": ("peaks.csv", csv_text.encode("utf-8"), "text/csv")})
    eq(r.status_code, 200, "CSV 上传应成功")
    d = r.json()
    eq(d["kind"], "peak_table", "用途推断")
    eq(d["count"], 2, "解析出的组分数")
    ok(d["file"]["sha256"], "必须记录内容哈希")
    r2 = c.post("/v1/imports/upload",
                files={"file": ("peaks.csv", csv_text.encode("utf-8"), "text/csv")})
    eq(r2.json()["file"]["reused"], True, "同内容应去重复用，不产生副本")
    a = c.post("/v1/analyses", json={"components": d["components"]}).json()
    eq(a["status"], "succeeded", "上传的峰表应能直接分析")


@case("API-13", "上传 EI 汇总表（xlsx）")
def _api13():
    if not Path(EI_TABLE).exists():
        print("        （跳过：EI 汇总表不可用）")
        return
    c = client()
    with open(EI_TABLE, "rb") as fh:
        r = c.post("/v1/imports/upload",
                   files={"file": ("PMAA汇总表.xlsx", fh,
                                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    eq(r.status_code, 200, "xlsx 上传应成功")
    d = r.json()
    eq(d["kind"], "ei_reference_table", "用途推断")
    eq(d["summary"]["total_rows"], 57, "应识别 57 行")
    eq(len(d["rows"]), 57, "逐行结果应完整返回")


@case("API-14", "拒收压缩包 / 未知类型 / 空文件")
def _api14():
    c = client()
    r = c.post("/v1/imports/upload",
               files={"file": ("data.zip", b"PK\x03\x04junk", "application/zip")})
    eq(r.status_code, 422, "压缩包应拒收")
    ok("解压" in r.json()["message"] or "展开" in r.json()["message"], "应说明不自动解压")
    # 注意：.pdf/.docx/图片现在都是受支持的格式，不能拿它们当"未知类型"
    r2 = c.post("/v1/imports/upload", files={"file": ("a.exe", b"MZ\x90\x00", "application/octet-stream")})
    eq(r2.status_code, 422, "未知类型应拒收")
    r3 = c.post("/v1/imports/upload", files={"file": ("empty.csv", b"", "text/csv")})
    eq(r3.status_code, 422, "空文件应拒收")


def _make_spectrum_image(truth: dict, *, mz_min: float, mz_max: float) -> bytes:
    """生成一张合成质谱棒图，用于验证图像反解精度。"""
    import io

    from PIL import Image, ImageDraw
    W, H, LEFT, RIGHT, BASE = 800, 400, 60, 780, 350
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    d.line([(LEFT, BASE), (RIGHT, BASE)], fill="black", width=2)
    for mz, inten in truth.items():
        x = LEFT + (mz - mz_min) / (mz_max - mz_min) * (RIGHT - LEFT)
        d.line([(x, BASE), (x, BASE - inten / 100.0 * 300)], fill="black", width=3)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _make_docx(truth: dict) -> bytes:
    import io

    import docx
    d = docx.Document()
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "m/z", "相对强度"
    for mz, i in sorted(truth.items()):
        row = t.add_row().cells
        row[0].text, row[1].text = str(mz), str(i)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _make_multi_image_docx() -> bytes:
    """生成含两张 EI 谱图与各自标签的 Word，用于多图回归。"""
    import io

    import docx
    d = docx.Document()
    for label, truth in (
        ("1,2,3,4-Galp", {87: 55, 118: 100, 129: 35}),
        ("1,2,3,4-Glcp", {59: 20, 102: 70, 162: 100}),
    ):
        d.add_paragraph(label)
        d.add_picture(io.BytesIO(_make_spectrum_image(truth, mz_min=40, mz_max=350)))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _make_pdf(truth: dict) -> bytes:
    import io

    import pymupdf as fitz
    doc = fitz.open()
    page = doc.new_page()
    lines = ["m/z   rel.int."] + [f"{mz}   {i}" for mz, i in sorted(truth.items())]
    page.insert_text((72, 72), "\n".join(lines), fontsize=11)
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


@case("DOC-01", "从图片反解质谱峰（m/z 与强度精度）")
def _doc01():
    from pmaa.ingest.spectrum_image import extract_peaks_from_image

    truth = {59: 12, 71: 20, 87: 55, 99: 51, 102: 50, 113: 70,
             118: 100, 129: 35, 145: 25, 162: 18, 233: 33}
    data = _make_spectrum_image(truth, mz_min=40.0, mz_max=350.0)
    res = extract_peaks_from_image(data, mz_min=40.0, mz_max=350.0)

    found = {round(p.mz): p.relative_intensity for p in res.peaks if p.mz is not None}
    eq([m for m in truth if m not in found], [], "不应有漏检的峰")
    eq([m for m in found if m not in truth], [], "不应有伪峰（坐标轴端点也不算峰）")
    errs = [abs(found[m] - truth[m]) for m in truth]
    ok(max(errs) < 2.0, f"强度最大误差 {max(errs):.2f} 应 < 2 个百分点")
    ok(res.confidence in ("low", "medium"), "置信度只能是 low/medium，永远不可能是 high")


@case("DOC-01b", "带标题、边框和文字的整页谱图只追踪与横轴相连的峰")
def _doc01b():
    import io

    from PIL import Image, ImageDraw
    from pmaa.ingest.spectrum_image import extract_peaks_from_image

    truth = {69: 20, 87: 80, 129: 100, 185: 55, 260: 25}
    w, h, left, right, baseline = 800, 600, 45, 745, 530
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    draw.text((280, 35), "PMAA spectral database - full spectrum", fill="black")
    draw.rectangle((20, 280, 770, 575), outline="black", width=2)
    draw.line((left, baseline, right, baseline), fill="black", width=2)
    for mz, intensity in truth.items():
        x = left + (mz - 50) / (340 - 50) * (right - left)
        draw.line((x, baseline, x, baseline - intensity * 2), fill="green", width=2)
        draw.text((x - 5, baseline - intensity * 2 - 18), str(mz), fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")

    res = extract_peaks_from_image(buf.getvalue(), mz_min=50, mz_max=340)
    found = [round(p.mz) for p in res.peaks if p.mz is not None]
    eq(res.params.get("extraction_mode"), "axis_attached", "应识别谱图横轴")
    eq(len(found), len(truth), "标题、数字标签和矩形边框不得变成伪峰")
    ok(all(abs(a - b) <= 1 for a, b in zip(found, sorted(truth))),
       "像素取整后的 m/z 误差应不超过 1")
    eq(len(found), len(set(found)), "同一根宽峰不得重复返回")


@case("DOC-02", "未给 m/z 轴范围时不得猜刻度")
def _doc02():
    from pmaa.ingest.spectrum_image import extract_peaks_from_image

    data = _make_spectrum_image({118: 100, 129: 35}, mz_min=40.0, mz_max=350.0)
    res = extract_peaks_from_image(data)                 # 故意不给轴范围
    ok(all(p.mz is None for p in res.peaks), "未校准时 m/z 必须为 None，不得臆造")
    ok(any("m/z" in w for w in res.warnings), "必须提示未给轴范围")
    ok(all(p.relative_intensity is not None for p in res.peaks), "强度仍应解出")


@case("DOC-03", "PDF / Word 提取：不得把强度列当成 m/z")
def _doc03():
    from pmaa.ingest.doc_parser import extract_docx, extract_pdf

    truth = {59: 12, 71: 20, 87: 55, 99: 51, 102: 50,
             113: 70, 118: 100, 129: 35, 145: 25, 162: 18, 233: 33}
    d = extract_docx(_make_docx(truth))
    eq(d.kind, "table", "应识别为表格路径")
    ions = d.components[0]["ions"]
    eq(len(ions), len(truth), "峰数应一致")
    ok(all(abs(i["intensity"] - truth[int(i["mz"])]) < 0.01 for i in ions), "强度应与原文一致")
    eq(d.components[0]["spectrum_completeness"], "selected_ions",
       "文档数值表不能伪装成仪器完整采集范围")

    # 关键：正文 "59  12" 里的 12 是强度，不能被当成 m/z
    p = extract_pdf(_make_pdf(truth))
    ok(p.kind in ("table", "text"), "应走文本/表格路径")
    mzs = {int(i["mz"]) for i in p.components[0]["ions"]}
    eq([m for m in mzs if m not in truth], [], "不得把强度值误收为 m/z")


@case("DOC-04", "提取接口：推导值必须被标注且证据上限降级")
def _doc04():
    c = client()
    data = _make_spectrum_image({118: 100, 129: 35, 87: 55}, mz_min=40.0, mz_max=350.0)
    r = c.post("/v1/imports/extract", files={"file": ("s.png", data, "image/png")},
               data={"mz_min": 40, "mz_max": 350, "run_analysis": "true"})
    eq(r.status_code, 200, "提取应成功")
    d = r.json()
    eq(d["needs_calibration"], True, "必须要求人工校准")
    eq(d["provenance"], "derived_from_image", "出处标注")
    ok(d.get("preview_image"), "应返回预览图供人工核对")
    a = (d.get("analysis") or [None])[0]
    ok(a is not None, "应顺带跑出候选")
    eq(a["derivative_identity_level"], "candidate", "推导值证据上限必须压到 candidate")
    ok("IMAGE_DERIVED" in a["missing_evidence"], "必须记录 IMAGE_DERIVED 缺失项")
    ok(any("非仪器直接导出" in n for n in a["notes"]), "必须给出推导值提示")

    r2 = c.post("/v1/imports/extract/confirm", json={
        "peaks": [{"mz": 118, "intensity": 100}, {"mz": 129, "intensity": 35}],
        "provenance": "derived_from_image"})
    eq(r2.status_code, 200, "确认应成功")
    eq(r2.json()["components"][0]["provenance"], "derived_from_image", "出处必须透传")
    eq(r2.json()["components"][0]["spectrum_completeness"], "selected_ions",
       "人工确认不能把推导峰表升级为完整仪器谱")
    r3 = c.post("/v1/analyses", json={
        "sample_reduction_label": "NaBD4",
        "components": r2.json()["components"],
    })
    eq(r3.status_code, 200, "网页真实链路：确认后的推导字段必须能进入分析 Schema")
    eq(r3.json()["components"][0]["result"]["derivative_identity_level"], "candidate",
       "推导谱进入分析后仍只能是候选")


@case("DOC-05", "老版 .doc 自动转换后进入内嵌谱图提取流程")
def _doc05():
    from pmaa.api import routers_core
    from pmaa.ingest.doc_parser import DocumentExtraction

    c = client()
    original = routers_core.extract_legacy_doc
    routers_core.extract_legacy_doc = lambda *_args, **_kwargs: DocumentExtraction(
        kind="image",
        images=[_make_spectrum_image({118: 100, 129: 35}, mz_min=40, mz_max=350)],
        meta={"converted_from": ".doc", "converter": "test"})
    try:
        r = c.post(
            "/v1/imports/upload",
            files={"file": ("old.doc", b"\xd0\xcf\x11\xe0test", "application/msword")},
            data={"mz_min": "40", "mz_max": "350"})
    finally:
        routers_core.extract_legacy_doc = original
    eq(r.status_code, 200, ".doc 应自动进入转换/提取流程")
    eq(r.json()["provenance"], "derived_from_document", "必须保留文档推导来源")
    ok(r.json().get("extracted_peaks"), "转换出的内嵌谱图应能继续提峰")


@case("DOC-05b", "损坏或伪装的 .doc 返回可理解错误而不是 500")
def _doc05b():
    c = client()
    r = c.post("/v1/imports/upload",
               files={"file": ("broken.doc", b"not a word document", "application/msword")})
    eq(r.status_code, 422, "损坏 .doc 应受控拒收")
    ok("不是可识别" in r.json()["message"], "应说明内容与扩展名不符")


@case("DOC-05c", "多图 Word：完整列图、逐图提取并批量固化为独立组分")
def _doc05c():
    c = client()
    word = _make_multi_image_docx()

    first = c.post(
        "/v1/imports/upload",
        files={"file": ("multi.docx", word,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"mz_min": "40", "mz_max": "350", "image_index": "0"})
    eq(first.status_code, 200, "多图 Word 应上传成功")
    d1 = first.json()
    eq(d1["embedded_images"], 2, "必须发现全部内嵌图片")
    eq(len(d1["image_inventory"]), 2, "必须返回完整图片清单")
    eq(d1["image_inventory"][0]["label"], "1,2,3,4-Galp", "第 1 张标签")
    eq(d1["image_inventory"][1]["label"], "1,2,3,4-Glcp", "第 2 张标签")

    second = c.post(
        "/v1/imports/extract",
        files={"file": ("multi.docx", word,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"mz_min": "50", "mz_max": "400", "image_index": "1"})
    eq(second.status_code, 200, "第 2 张应可独立提取")
    d2 = second.json()
    eq(d2["selected_image_index"], 1, "应返回第 2 张")
    eq(d2["selected_image_label"], "1,2,3,4-Glcp", "所选图片标签应正确")

    confirmed = c.post("/v1/imports/extract/confirm", json={"spectra": [
        {"image_index": 0, "label": "1,2,3,4-Galp · 图 1",
         "peaks": [{"mz": 87, "marker": "+"}, {"mz": 118, "marker": "red"}],
         "qualitative": True, "provenance": "derived_from_document",
         "extraction_params": {"mz_min": 40, "mz_max": 350}},
        {"image_index": 1, "label": "1,2,3,4-Glcp · 图 2",
         "peaks": [{"mz": 102, "marker": "red"}, {"mz": 162, "marker": "red"}],
         "qualitative": True, "provenance": "derived_from_document",
         "extraction_params": {"mz_min": 50, "mz_max": 400}},
    ]})
    eq(confirmed.status_code, 200, "多谱批量确认应成功")
    result = confirmed.json()
    eq(result["count"], 2, "两张谱图必须生成两个组分")
    eq([x["component_id"] for x in result["components"]],
       ["extracted_image_001", "extracted_image_002"], "组分 ID 应逐图唯一")
    eq(result["components"][0]["extraction_params"]["mz_min"], 40,
       "第 1 张独立轴范围必须保留")
    eq(result["components"][1]["extraction_params"]["mz_min"], 50,
       "第 2 张独立轴范围必须保留")

    duplicate = c.post("/v1/imports/extract/confirm", json={"spectra": [
        {"image_index": 0, "peaks": [{"mz": 118, "marker": "+"}]},
        {"image_index": 0, "peaks": [{"mz": 129, "marker": "+"}]},
    ]})
    eq(duplicate.status_code, 422, "重复图片序号必须拒绝，不能覆盖或合并")
    eq(duplicate.json()["error_code"], "DUPLICATE_IMAGE_INDEX", "稳定错误码")


@case("DOC-05d", "多图界面必须提供上一张/下一张按钮与缓存更新标识")
def _doc05d():
    c = client()
    page = c.get("/")
    script = c.get("/app.js")
    eq(page.status_code, 200, "平台首页应可读取")
    ok("simple-3step-1" in page.text, "app.js 必须带缓存版本，避免浏览器继续使用旧界面")
    eq(script.status_code, 200, "前端脚本应可读取")
    ok("上传文件" in script.text and "检查谱图与峰" in script.text and "查看结果" in script.text,
       "主流程必须是三步简化版")
    ok("runSimplifiedAnalysis" in script.text and "go('import')" in script.text,
       "确认峰后必须自动解析，并从上传页启动")
    ok("LLM 配置" in page.text, "LLM 配置必须保留在顶部")
    for token in ('id="x_prev"', 'id="x_next"', "async function switchImage", "下一张 →"):
        ok(token in script.text, f"多图切换界面必须包含 {token}")


@case("DOC-07", "OCR 自动识别 x 轴刻度（未装 OCR 则跳过）")
def _doc07():
    """
    若装有 rapidocr，上传带刻度标签的谱图且不填轴范围时，
    应自动标定出 m/z 轴，且应接近真值。
    未装 OCR 的环境直接跳过——该能力本就是可选的。
    """
    try:
        from pmaa.ingest.ocr_axis import ocr_available
        if not ocr_available():
            print("        （跳过：未安装 rapidocr）")
            return
    except Exception:
        print("        （跳过：ocr_axis 不可用）")
        return

    c = client()
    data = _make_spectrum_image_with_axis(mz_min=40.0, mz_max=350.0)
    r = c.post("/v1/imports/upload", files={"file": ("axis.png", data, "image/png")})
    eq(r.status_code, 200, "应上传成功")
    d = r.json()
    ac = d.get("axis_calibration") or {}
    ok(ac.get("ok"), f"应自动标定成功，实际 {ac}")
    ok(ac.get("n_points", 0) >= 3, "应识别到 ≥3 个刻度")
    ok((ac.get("r2") or 0) >= 0.95, "拟合 R² 应达标")
    ok(abs(ac["mz_min"] - 40.0) <= 3.0, f"轴起点应接近 40，实际 {ac['mz_min']}")
    ok(abs(ac["mz_max"] - 350.0) <= 3.0, f"轴终点应接近 350，实际 {ac['mz_max']}")
    ok(all(p["mz"] is not None for p in d["extracted_peaks"]), "m/z 应全部解出")


@case("DOC-06", "上传入口必须自动路由图片/PDF/Word，不能报'不支持的峰表格式'")
def _doc06():
    """
    回归用例：曾经 /imports/upload 只处理 ei_reference_table，其余一律掉进
    parse_peak_table()，导致 PNG 被报 422「不支持的峰表格式」——
    而平台明明在格式列表里宣称支持 PNG，用户就会以为传不上去。
    """
    c = client()
    data = _make_spectrum_image({118: 100, 129: 35, 87: 55}, mz_min=40.0, mz_max=350.0)

    r = c.post("/v1/imports/upload", files={"file": ("s.png", data, "image/png")})
    eq(r.status_code, 200, "PNG 上传不得再报 422")
    d = r.json()
    eq(d["needs_calibration"], True, "应转入需人工校准的提取结果")
    ok(d.get("preview_image"), "应带回预览图")

    r2 = c.post("/v1/imports/upload", files={"file": ("s.png", data, "image/png")},
                data={"mz_min": "40", "mz_max": "350"})
    eq(r2.status_code, 200, "带轴范围应成功")
    d2 = r2.json()
    eq(d2["calibrated"], True, "应已标定")
    ok(all(p["mz"] is not None for p in d2["extracted_peaks"]), "m/z 应全部解出")

    r3 = c.post("/v1/imports/upload",
                files={"file": ("s.pdf", _make_pdf({118: 100}), "application/pdf")})
    eq(r3.status_code, 200, "PDF 不得报 422")
    r4 = c.post("/v1/imports/upload",
                files={"file": ("s.docx", _make_docx({118: 100}),
                                "application/vnd.openxmlformats-officedocument."
                                "wordprocessingml.document")})
    eq(r4.status_code, 200, "DOCX 不得报 422")

    # 结构化文件仍走原路径，未被误伤
    r5 = c.post("/v1/imports/upload",
                files={"file": ("p.csv", b"component_id,evidence_type,ions\n"
                                         b"p1,qualitative_ion_list,\"118:+\"\n", "text/csv")})
    eq(r5.json()["kind"], "peak_table", "CSV 仍走峰表路径")


@case("API-15", "N1–N10 制备记录绑定分析并进入 LLM 证据包")
def _api15():
    from pmaa.data import repositories as repo

    c = client()
    c.post("/v1/samples", json={"sample_id": "prep_sample", "code": "PREP-1"})
    prep = c.post("/v1/samples/prep_sample/preparations", json={
        "preparation_batch_id": "prep_batch_1", "batch_code": "PB1",
        "sop_version": "neutral_N_v1",
        "independent_preparation_count": 1, "post_methylation_aliquots": 4,
        "steps": {
            "N1": {"status": "completed", "actual_values": {"sample_mg": 10.2}},
            "N7": {"status": "completed", "actual_values": {"reagent": "NaBD4"}},
        }}).json()
    eq(prep["steps"]["sop_version"], "neutral_N_v1", "SOP 版本")
    aid = c.post("/v1/analyses", json={
        "sample_id": "prep_sample", "preparation_batch_id": "prep_batch_1",
        "sample_reduction_label": "NaBD4",
        "components": [{"component_id": "prep_peak", "evidence_type": "qualitative_ion_list",
                        "ions": [{"mz": 118, "marker": "+"}]}]}).json()["analysis_id"]
    assessed = c.post(f"/v1/analyses/{aid}/llm-assessments", json={
        "component_ids": ["prep_peak"], "provider_config_id": "api_mock"}).json()
    bid = assessed["assessments"][0]["bundle_id"]
    fields = {f["field"] for f in repo.BundleRepo.get(bid).payload["facts"]}
    ok("preparation.N1" in fields and "preparation.N7" in fields,
       "实际步骤记录必须进入服务端冻结证据包")


@case("API-16", "相同 component_id 的不同分析不得串入 LLM 评议")
def _api16():
    c = client()
    aid = c.post("/v1/analyses", json={
        "components": [{"component_id": "peak_007", "evidence_type": "qualitative_ion_list",
                        "ions": [{"mz": 118, "marker": "+"}]}]}).json()["analysis_id"]
    listed = c.get(f"/v1/analyses/{aid}/llm-assessments").json()["items"]
    eq(listed, [], "不得按 component_id 跨分析串数据")
    rep = c.post(f"/v1/analyses/{aid}/reports", json={"fmt": "json"}).json()
    eq(rep["llm"]["enabled"], False, "报告不得串入其他分析的 LLM 评议")
    eq(rep["components"][0]["review_history"], [], "人工决定也不得按 component_id 跨分析串入")


def _make_spectrum_image_with_axis(truth: dict | None = None, *, mz_min: float,
                                   mz_max: float) -> bytes:
    """生成一张**带 x 轴刻度数字**的谱图，用于验证 OCR 轴校准。"""
    import io

    from PIL import Image, ImageDraw, ImageFont
    truth = truth or {59: 12, 71: 20, 87: 55, 99: 51, 102: 50, 113: 70,
                      118: 100, 129: 35, 145: 25, 162: 18, 233: 33}
    W, H, LEFT, RIGHT, BASE = 900, 460, 80, 860, 380
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 16)
    except Exception:
        font = ImageFont.load_default()

    d.line([(LEFT, BASE), (RIGHT, BASE)], fill="black", width=2)
    for mz, inten in truth.items():
        x = LEFT + (mz - mz_min) / (mz_max - mz_min) * (RIGHT - LEFT)
        d.line([(x, BASE), (x, BASE - inten / 100.0 * 300)], fill="black", width=3)
    # x 轴刻度 + 数字标签（OCR 要读的就是这些）
    start = int(mz_min // 50 * 50) + 50
    for mz in range(start, int(mz_max) + 1, 50):
        x = LEFT + (mz - mz_min) / (mz_max - mz_min) * (RIGHT - LEFT)
        d.line([(x, BASE), (x, BASE + 6)], fill="black", width=1)
        label = str(mz)
        bb = d.textbbox((0, 0), label, font=font)
        d.text((x - (bb[2] - bb[0]) / 2, BASE + 10), label, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()
