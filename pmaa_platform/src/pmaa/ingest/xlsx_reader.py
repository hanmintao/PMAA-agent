"""
XLSX 只读解析（M02 表格导入）。

设计要点：
* 只读打开，不执行宏、不写回原文件；
* 每个取值都带 source_cell（如 Sheet1!R3C12），可追溯到源文件单元格（V0.1 完成条件）；
* 识别合并单元格、字体颜色、填充色、隐藏行/列、说明文字块；
* 颜色语义**不换算成数值强度**——红色只记为 reported_gt20pct（14.8 第 3 步 / A01）。
"""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

RED_RGB_SUFFIXES = ("FFFF0000", "00FF0000", "FF0000")


@dataclass
class Cell:
    coord: str                 # 例 "C3"
    row: int
    col: int
    value: Any
    font_color: str | None = None
    fill_color: str | None = None
    bold: bool = False
    is_merged: bool = False
    merged_range: str | None = None

    @property
    def is_red(self) -> bool:
        if not self.font_color:
            return False
        c = self.font_color.upper()
        return any(c.endswith(s) for s in RED_RGB_SUFFIXES)


@dataclass
class SheetDump:
    name: str
    max_row: int
    max_col: int
    merged_ranges: list[str] = field(default_factory=list)
    hidden_rows: list[int] = field(default_factory=list)
    hidden_cols: list[int] = field(default_factory=list)
    cells: dict[tuple[int, int], Cell] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)

    def cell(self, row: int, col: int) -> Cell | None:
        return self.cells.get((row, col))

    def row_values(self, row: int, col_from: int, col_to: int) -> list[Cell | None]:
        return [self.cell(row, c) for c in range(col_from, col_to + 1)]


def _color_str(color) -> str | None:
    if color is None:
        return None
    try:
        if getattr(color, "rgb", None) and isinstance(color.rgb, str):
            return color.rgb
        if getattr(color, "theme", None) is not None:
            return f"theme:{color.theme}"
        if getattr(color, "indexed", None) is not None:
            return f"indexed:{color.indexed}"
    except Exception:
        return None
    return None


def read_workbook(path: str | Path) -> tuple[dict[str, SheetDump], str, int]:
    """
    只读解析整个工作簿。

    :return: (sheets, sha256, byte_size)
    """
    p = Path(path)
    raw = p.read_bytes()
    sha = hashlib.sha256(raw).hexdigest()

    wb = openpyxl.load_workbook(p, data_only=False, read_only=False)
    sheets: dict[str, SheetDump] = {}
    for ws in wb.worksheets:
        dump = SheetDump(name=ws.title, max_row=ws.max_row or 0, max_col=ws.max_column or 0)
        dump.merged_ranges = [str(r) for r in ws.merged_cells.ranges]

        # 隐藏行列：导入预览时必须提示，避免"看不见的数据"被静默忽略
        for r_idx, dim in getattr(ws, "row_dimensions", {}).items():
            if getattr(dim, "hidden", False):
                dump.hidden_rows.append(r_idx)
        for c_letter, dim in getattr(ws, "column_dimensions", {}).items():
            if getattr(dim, "hidden", False):
                col_idx = openpyxl.utils.column_index_from_string(c_letter)
                dump.hidden_cols.append(col_idx)

        merged_map = {}
        for rng in ws.merged_cells.ranges:
            for row in ws[rng.coord]:
                for c in row:
                    merged_map[(c.row, c.column)] = str(rng)

        for row in ws.iter_rows():
            for c in row:
                if c.value is None or (isinstance(c.value, str) and not c.value.strip()):
                    continue
                dump.cells[(c.row, c.column)] = Cell(
                    coord=f"{get_column_letter(c.column)}{c.row}",
                    row=c.row, col=c.column, value=c.value,
                    font_color=_color_str(c.font.color) if c.font else None,
                    fill_color=_color_str(c.fill.fgColor) if c.fill else None,
                    bold=bool(c.font and c.font.bold),
                    is_merged=(c.row, c.column) in merged_map,
                    merged_range=merged_map.get((c.row, c.column)),
                )
        sheets[ws.title] = dump
    wb.close()
    return sheets, sha, len(raw)


# ---------------------------------------------------------------------- 汇总表专用结构识别
@dataclass
class SummaryTableLayout:
    """针对『PMAA的EI质谱图的数据汇总表』这类"行=衍生物、列=m/z"的宽表布局识别结果。"""
    sheet: str
    header_row: int
    name_col: int
    linkage_col: int | None
    ion_columns: dict[int, float]      # col_index -> m/z
    data_rows: list[int]
    section_rows: dict[int, str]       # row -> 分组标题（如 六碳糖）
    note_rows: list[int]
    detected_semantics: dict[str, Any] = field(default_factory=dict)


_SECTION_NAME_RE = re.compile(r"(六碳糖|五碳糖|脱氧|6-脱氧己糖)")


def detect_summary_layout(dump: SheetDump, sheet_name: str | None = None) -> SummaryTableLayout:
    """
    自动识别布局：
      * 表头行 = 第一个含多个纯数字单元格（m/z）的行；
      * 名称列 = 第 1 列（含 '-' 且非纯数字）；
      * 连接方式列 = 表头为 "连接方式" 的列；
      * 分组行 = 只有一个非空单元格且匹配 六碳糖/五碳糖/脱氧；
      * 注释行 = 合并单元格中的长文本。
    """
    header_row, best = 1, 0
    for r in range(1, min(dump.max_row, 10) + 1):
        nums = 0
        for c in range(1, dump.max_col + 1):
            cell = dump.cell(r, c)
            if cell is not None and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                nums += 1
        if nums > best:
            header_row, best = r, nums

    ion_columns: dict[int, float] = {}
    linkage_col = None
    for c in range(1, dump.max_col + 1):
        cell = dump.cell(header_row, c)
        if cell is None:
            continue
        if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
            ion_columns[c] = float(cell.value)
        elif isinstance(cell.value, str) and "连接" in cell.value:
            linkage_col = c

    section_rows: dict[int, str] = {}
    note_rows: list[int] = []
    data_rows: list[int] = []

    for r in range(header_row + 1, dump.max_row + 1):
        cells = [dump.cell(r, c) for c in range(1, dump.max_col + 1)]
        non_empty = [c for c in cells if c is not None]
        if not non_empty:
            continue
        first = non_empty[0]
        first_val = str(first.value)
        # 先判注释：合并单元格中的长文本块（常含"红色…20%""NaBD4"等语义声明），
        # 必须先于分组标题判断，否则含"五碳糖"字样的注释会被误判成分组行。
        if (len(non_empty) == 1 and first.is_merged and len(first_val) > 20) or \
           (len(non_empty) == 1 and first.is_merged and any(
               kw in first_val for kw in ("红色", "NaBD4", "NaBH4", "信号"))):
            note_rows.append(r)
            continue
        if len(non_empty) == 1 and not first.is_merged and len(first_val) <= 12 \
                and _SECTION_NAME_RE.search(first_val):
            section_rows[r] = first_val.strip()
            continue
        # 数据行：名称列必须像 PMAA 名（含 '-'）
        name_cell = dump.cell(r, 1)
        if name_cell is not None and isinstance(name_cell.value, str) and "-" in name_cell.value:
            data_rows.append(r)

    # 从注释行提取语义声明（不猜，只记录原文匹配到的声明）
    semantics: dict[str, Any] = {}
    for r in note_rows:
        # 合并单元格的取值可能落在任意列，取该行第一个非空单元格
        note_cell = next((dump.cell(r, c) for c in range(1, dump.max_col + 1)
                          if dump.cell(r, c) is not None), None)
        txt = str(note_cell.value) if note_cell else ""
        if "红色" in txt and "20%" in txt:
            semantics["red_means"] = "reported_gt20pct"
        if "+" in txt and "有该信号" in txt:
            semantics["plus_means"] = "reported_present"
        if "NaBD4" in txt:
            m = re.search(r"(NaBD4|NaBH4)", txt)
            if m and "而不是" in txt:
                semantics["reduction_label"] = m.group(1)
            elif m:
                semantics.setdefault("reduction_label", m.group(1))
        if "呋喃" in txt:
            semantics["pentose_ring_note"] = txt.strip()
    semantics["detected_merged_note_rows"] = note_rows

    return SummaryTableLayout(
        sheet=sheet_name or dump.name, header_row=header_row, name_col=1,
        linkage_col=linkage_col, ion_columns=ion_columns, data_rows=data_rows,
        section_rows=section_rows, note_rows=note_rows, detected_semantics=semantics,
    )
